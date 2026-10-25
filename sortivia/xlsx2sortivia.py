"""Convierte una hoja de claves ESTOPA (formato v2) en una actividad de SORTIVIA.

Pestañas del Excel:
  Actividad  → parámetros: id, títulos y resúmenes (es/en), presupuesto, rondas, números del sorteo,
               puntos por número, puntos_objetivo (puntos de la mejor estrategia posible) y bloques obligatorios
  Bloques    → Bloque (EN) | Bloque (ES) | Objetivo (EN) | Objetivo (ES)
               Los bloques con el mismo objetivo puntúan una sola vez (cuenta el mejor).
  Opciones   → Bloque | Decisión | Grupo | Opción | Tipo de coste | Precio | Puntos | Notas | Selección |
               Decisión (ES) | Opción (ES) | Grupo (ES)
               Tipos de coste: réplicas, por muestra, multiplicador, fijo, por unidad, condiciones
               Selección: única o múltiple. En las múltiples cuenta la mejor opción positiva marcada y restan todas las negativas.
  Reglas     → ID | Condición | Puntos | Explicación (ES) | Explicación (EN) | Bloque (opcional)   (ocultas a los alumnos)
               Con «Bloque», los puntos se suman a ese bloque y solo cuentan si es el que puntúa en su objetivo.

Genera:
  actividades/<id>.json      → pública (sin puntos ni reglas): va al repositorio
  claves/<id>.clave.json     → privada (puntos, reglas y escala): se carga en el panel, NUNCA en el repositorio
Calcula la mejor estrategia posible dentro del presupuesto y escala los puntos para que obtenga «puntos_objetivo».
Escribe en el Excel las pestañas «Identificadores» y «Estrategia óptima».

Uso: python3 -I xlsx2sortivia.py claves_v2.xlsx carpeta_actividades/ carpeta_claves/ [--rapido]
"""
import json, os, re, sys, unicodedata
import openpyxl
from openpyxl.styles import Font
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import sortivia_opt as OPT

args = [a for a in sys.argv[1:] if not a.startswith("--")]
src, out_pub_dir, out_key_dir = args
RAPIDO = "--rapido" in sys.argv
TIPOS = {"réplicas": "replicas", "replicas": "replicas", "por muestra": "muestra", "multiplicador": "mult", "fijo": "fijo", "por unidad": "unidad", "condiciones": "condiciones"}

def slug(s):
    # Lo que va entre corchetes (p. ej., «de novo [SPAdes, Flye]») no forma parte del identificador:
    # se pueden actualizar los programas sin romper las reglas.
    s = re.sub(r"\[[^\]]*\]", "", str(s)).strip() or str(s)
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")
def num(v):
    if isinstance(v, (int, float)): return float(v)
    return float(str(v).strip().lower().lstrip("x").replace(",", ".") or 0)
def txt(en, es): en = str(en).strip() if en is not None else ""; es = str(es).strip() if es else ""; return {"en": en, "es": es or en}

wb = openpyxl.load_workbook(src)
P = {str(r[0]).strip(): r[1] for r in wb["Actividad"].iter_rows(min_row=2, values_only=True) if r[0]}
aid = slug(P["id"])
obligatorios = [slug(x) for x in str(P.get("bloques_obligatorios") or "").split(";") if x.strip()]

# Bloques: nombre en los dos idiomas y objetivo
info_bloques = {}
if "Bloques" in wb.sheetnames:
    for row in wb["Bloques"].iter_rows(min_row=2, values_only=True):
        if not row or not row[0] or not row[1] and not row[2]: continue
        en, es, oen, oes = (list(row) + [None] * 4)[:4]
        info_bloques[slug(en)] = {"nombre": txt(en, es), "objetivo": {"id": slug(oen), "nombre": txt(oen, oes)} if oen else None}

bloques, idx, puntos, errores = [], {}, {}, []
for i, row in enumerate(wb["Opciones"].iter_rows(min_row=2, values_only=True), start=2):
    if not row or not row[0]: continue
    B, D, G, O, TIPO, PRECIO, PTS, _NOTAS, SELEC, D_ES, O_ES, G_ES = (list(row) + [None] * 12)[:12]
    tipo = TIPOS.get(str(TIPO).strip().lower())
    if not tipo: errores.append(f"Opciones fila {i}: tipo de coste desconocido «{TIPO}»"); continue
    bid, did = slug(B), slug(D)
    if bid not in idx:
        ib = info_bloques.get(bid, {})
        idx[bid] = {"id": bid, "nombre": ib.get("nombre") or txt(B, None), "decisiones": [], "_d": {}}
        if ib.get("objetivo"): idx[bid]["objetivo"] = ib["objetivo"]
        bloques.append(idx[bid])
    blk = idx[bid]
    if did not in blk["_d"]:
        multiple = str(SELEC or "").strip().lower().startswith("m")
        dec = {"id": did, "nombre": txt(D, D_ES), "tipo": "contador" if tipo == "unidad" else ("multiple" if multiple else "unica"), "opciones": []}
        blk["_d"][did] = dec; blk["decisiones"].append(dec)
    dec = blk["_d"][did]
    oid = "n" if tipo == "unidad" else slug(O)
    coste = {"tipo": tipo, "valor": num(PRECIO or 0)}
    if tipo == "replicas": coste["n"] = int(num(O))
    opt = {"id": oid, "label": txt(O, O_ES), "coste": coste}
    if G: opt["grupo"] = txt(G, G_ES)
    if any(o["id"] == oid for o in dec["opciones"]): errores.append(f"Opciones fila {i}: opción repetida «{O}» en {B} · {D}")
    dec["opciones"].append(opt)
    puntos[f"{bid}/{did}/{oid}"] = num(PTS or 0)
for b in bloques: del b["_d"]

# Reglas de coherencia
TERM = re.compile(r"^(!?)([a-z0-9-]+)(?:/([a-z0-9-]+)(?:(=)([a-z0-9|-]+)|(>)(\d+)))?$")
reglas = []
if "Reglas" in wb.sheetnames:
    for i, row in enumerate(wb["Reglas"].iter_rows(min_row=2, values_only=True), start=2):
        if not row or not row[0] or not row[1]: continue
        rid, cond, pts, t_es, t_en, r_blk = (list(row) + [None] * 6)[:6]
        terms = []
        for t in [x.strip() for x in str(cond).split("&") if x.strip()]:
            m = TERM.match(t)
            if not m: errores.append(f"Reglas fila {i} ({rid}): no entiendo «{t}»"); continue
            neg, b, d, eq, opts, gt, n = m.groups()
            if b not in idx: errores.append(f"Reglas fila {i} ({rid}): no existe el bloque «{b}»"); continue
            term = {"bloque": b, "no": bool(neg)}
            if d:
                dec = next((x for x in idx[b]["decisiones"] if x["id"] == d), None)
                if not dec: errores.append(f"Reglas fila {i} ({rid}): no existe la decisión «{b}/{d}»"); continue
                term["decision"] = d
                if eq:
                    ids = opts.split("|")
                    for o in ids:
                        if not any(x["id"] == o for x in dec["opciones"]): errores.append(f"Reglas fila {i} ({rid}): no existe la opción «{b}/{d}={o}»")
                    term["opciones"] = ids
                else:
                    term["mayor_que"] = int(n)
            terms.append(term)
        regla = {"id": str(rid).strip(), "terminos": terms, "puntos": num(pts or 0), "texto": txt(t_en or t_es, t_es)}
        if r_blk:
            if slug(r_blk) not in idx: errores.append(f"Reglas fila {i} ({rid}): no existe el bloque «{r_blk}» de la columna Bloque")
            else: regla["bloque"] = slug(r_blk)
        reglas.append(regla)

if errores:
    print("ERRORES:"); [print("  -", e) for e in errores]; sys.exit(1)

act = {
    "schema": "sortivia-actividad/3", "id": aid, "metodologia": "ESTOPA",
    "titulo": {"es": P.get("titulo_es") or P["id"], "en": P.get("titulo_en") or P.get("titulo_es") or P["id"]},
    "resumen": {"es": P.get("resumen_es") or "", "en": P.get("resumen_en") or P.get("resumen_es") or ""},
    "presupuesto": num(P.get("presupuesto") or 0), "numeros": int(num(P.get("numeros") or 10)), "rondas": int(num(P.get("rondas") or 2)),
    "puntos_por_numero": num(P.get("puntos_por_numero") or 10), "bloques_obligatorios": obligatorios, "bloques": bloques,
}
if P.get("unidad_es") or P.get("unidad_en"):   # otra unidad para el «presupuesto» (p. ej., días)
    act["unidad"] = {"es": str(P.get("unidad_es") or P.get("unidad_en")).strip(), "en": str(P.get("unidad_en") or P.get("unidad_es")).strip()}
if P.get("asignatura"):   # solo para organizar las actividades en el panel
    act["asignatura"] = str(P.get("asignatura")).strip()
clave = {"schema": "sortivia-clave/3", "actividad": aid, "puntos": puntos, "reglas": reglas}

# Estrategia óptima dentro del presupuesto → escala de puntos
objetivo = num(P.get("puntos_objetivo") or 99)
mejores = []
for semilla in ((1, 2) if RAPIDO else (1, 2, 3, 4)):
    mejores.append(OPT.optimizar(act, clave, act["presupuesto"], reinicios=8 if RAPIDO else 14, iteraciones=5000 if RAPIDO else 9000, semilla=semilla))
bruto, plan = max(mejores, key=lambda x: x[0])
clave["escala"] = {"objetivo": objetivo, "maximo": bruto, "presupuesto": act["presupuesto"]}

os.makedirs(out_pub_dir, exist_ok=True); os.makedirs(out_key_dir, exist_ok=True)
json.dump(act, open(os.path.join(out_pub_dir, f"{aid}.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
json.dump(clave, open(os.path.join(out_key_dir, f"{aid}.clave.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)

# Pestaña de identificadores para escribir reglas
if "Identificadores" in wb.sheetnames: del wb["Identificadores"]
wi = wb.create_sheet("Identificadores")
wi.append(["Bloque", "Decisión", "Opción", "Para usar en una regla"])
for b in bloques:
    for d in b["decisiones"]:
        for o in d["opciones"]:
            wi.append([b["nombre"]["en"], d["nombre"]["en"], o["label"]["en"], f"{b['id']}/{d['id']}>0" if d["tipo"] == "contador" else f"{b['id']}/{d['id']}={o['id']}"])
for c, w in zip("ABCD", (20, 24, 46, 80)): wi.column_dimensions[c].width = w

# Pestaña con la estrategia óptima encontrada
if "Estrategia óptima" in wb.sheetnames: del wb["Estrategia óptima"]
wo = wb.create_sheet("Estrategia óptima")
total_coste = OPT.coste_total(act, plan)
_, det = OPT.puntos_detalle(act, clave, plan)
U = act["unidad"]["es"].split("|")[-1] if act.get("unidad") else "€"
wo.append([f"Mejor estrategia encontrada con {act['presupuesto']:.0f} {U}: {bruto:.0f} puntos brutos → {objetivo:.0f} puntos ({int(objetivo // act['puntos_por_numero'])} números). Coste: {total_coste:.0f} {U}."])
wo["A1"].font = Font(bold=True)
wo.append(["Es una referencia calculada automáticamente; si ves algo que no tiene sentido, ajusta los puntos o las reglas y vuelve a convertir."])
wo.append([]); wo.append(["Bloque", "Decisión", "Elección", "Puntos que cuentan", f"Coste del bloque ({U})"])
for c in wo[4]: c.font = Font(bold=True)
nc = OPT.condiciones(act, plan); nr = OPT.replicas(act, plan)
for b in bloques:
    if not OPT.incluido(act, plan, b): continue
    cb = OPT.coste_bloque(b, plan[b["id"]], nc, nr)
    for d in b["decisiones"]:
        v = plan[b["id"]]["dec"].get(d["id"])
        if d["tipo"] == "contador":
            if int(v or 0) <= 0: continue
            elec = f"{int(v)} ×"
        else:
            elec = ", ".join(o["label"]["en"] for o in OPT.elegidas(d, v))
        pc = sum(p for bb, dd, oo, p, c in det if bb == b["id"] and dd == d["id"] and c)
        wo.append([b["nombre"]["en"], d["nombre"]["en"], elec, pc, cb])
reg = [(dd, p) for bb, dd, oo, p, c in det if bb == "regla"]
if reg:
    wo.append([]); wo.append(["Reglas que se cumplen", ", ".join(f"{r} ({p:+.0f})" for r, p in reg)])
for c, w in zip("ABCDE", (20, 30, 60, 18, 20)): wo.column_dimensions[c].width = w
wb.save(src)
print(f"{aid}: {len(bloques)} bloques, {sum(len(b['decisiones']) for b in bloques)} decisiones, {len(puntos)} opciones, {len(reglas)} reglas")
print(f"Mejor estrategia con {act['presupuesto']:.0f} {U}: {bruto:.0f} puntos brutos (coste {total_coste:.0f} {U}) → escala a {objetivo:.0f} puntos")
